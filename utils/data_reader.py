"""
data_reader.py: Test Data Reader Utility for Capstone V1.
Supports loading and updating test configuration and credentials from either JSON or Excel (.xlsx).
"""

import json
import os
import openpyxl


def load_json_data(file_path: str) -> dict:
    """Load test data from a JSON file."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"JSON test data file not found: {file_path}")
    with open(file_path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_json_data(file_path: str, data: dict) -> None:
    """Save test data to a JSON file."""
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)


def load_excel_data(file_path: str) -> dict:
    """Load test data from an Excel (.xlsx) file containing UserCredentials and ProductData sheets."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Excel test data file not found: {file_path}")
    
    wb = openpyxl.load_workbook(file_path)
    data = {
        "base_url": "https://tutorialsninja.com/demo/",
        "user": {},
        "product": {}
    }
    
    # Read UserCredentials sheet
    if "UserCredentials" in wb.sheetnames:
        ws_user = wb["UserCredentials"]
        headers = [cell.value for cell in ws_user[1]]
        values = [cell.value for cell in ws_user[2]]
        for h, v in zip(headers, values):
            if h is not None:
                data["user"][str(h).strip()] = v
                
    # Read ProductData sheet
    if "ProductData" in wb.sheetnames:
        ws_prod = wb["ProductData"]
        headers = [cell.value for cell in ws_prod[1]]
        values = [cell.value for cell in ws_prod[2]]
        for h, v in zip(headers, values):
            if h is not None:
                data["product"][str(h).strip()] = v
                
    wb.close()
    return data


def save_excel_user(file_path: str, user_data: dict) -> None:
    """Update user credentials in the Excel file."""
    if not os.path.exists(file_path):
        return
    wb = openpyxl.load_workbook(file_path)
    if "UserCredentials" in wb.sheetnames:
        ws = wb["UserCredentials"]
        headers = [str(cell.value).strip() for cell in ws[1] if cell.value is not None]
        for col_idx, header in enumerate(headers, start=1):
            if header in user_data:
                ws.cell(row=2, column=col_idx, value=user_data[header])
        wb.save(file_path)
    wb.close()


def get_test_data(source_type: str = "json", base_dir: str = None) -> dict:
    """
    Unified entry point to fetch test data.
    :param source_type: 'json' or 'excel'
    :param base_dir: directory containing data.json / data.xlsx
    """
    if base_dir is None:
        base_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "test_data")
        
    json_path = os.path.join(base_dir, "data.json")
    excel_path = os.path.join(base_dir, "data.xlsx")
    
    if source_type.lower() == "excel":
        return load_excel_data(excel_path)
    else:
        return load_json_data(json_path)


def persist_user_credentials(user_data: dict, base_dir: str = None) -> None:
    """Persist updated user credentials to both JSON and Excel files."""
    if base_dir is None:
        base_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "test_data")
        
    json_path = os.path.join(base_dir, "data.json")
    excel_path = os.path.join(base_dir, "data.xlsx")
    
    if os.path.exists(json_path):
        data = load_json_data(json_path)
        data["user"].update(user_data)
        save_json_data(json_path, data)
        
    if os.path.exists(excel_path):
        save_excel_user(excel_path, user_data)
